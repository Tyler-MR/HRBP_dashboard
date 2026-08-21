# -*- coding: utf-8 -*-
"""
钉钉日志 → Excel 导出
读取 raw_logs.json，生成多 sheet Excel：
  1. 总览表：全部日志（日期/人员/部门/模板/内容摘要）
  2. 各模板明细 sheet：字段为列
"""
import json, time
from datetime import datetime
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

RAW = Path.home() / 'Desktop' / '日志导出' / 'raw_logs.json'
OUT = Path.home() / 'Desktop' / '日志导出' / '钉钉日志导出.xlsx'

INDIGO = '4F46E5'   # 主题色-靛蓝
GREEN = '059669'    # 主题色-翠绿
LIGHT = 'EEF2FF'    # 浅靛蓝底

import re as _re
_ILLEGAL = _re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')

def clean(v):
    """过滤 Excel 非法控制字符（钉钉内容含 \b 等）。"""
    if isinstance(v, str):
        return _ILLEGAL.sub('', v)
    return v

def fmt_ts(ms):
    return datetime.fromtimestamp(ms / 1000).strftime('%Y-%m-%d %H:%M')

def content_text(x):
    parts = []
    for c in (x.get('contents') or []):
        v = clean((c.get('value') or '')).strip()
        if v:
            parts.append(v)
    return '\n'.join(parts)

def style_header(ws, ncols, title=None):
    if title:
        ws.insert_rows(1)
        ws.cell(row=1, column=1, value=title).font = Font(size=14, bold=True, color='FFFFFF')
        ws.cell(row=1, column=1).fill = PatternFill('solid', fgColor=INDIGO)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
        ws.cell(row=1, column=1).alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 24
    header_row = 2 if title else 1
    for c in range(1, ncols + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=INDIGO)
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

def auto_width(ws, max_w=60):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max(len(str(c.value)) for c in col[:60] if c.value) if any(c.value for c in col[:60]) else 10
        ws.column_dimensions[letter].width = min(max(width * 1.4, 10), max_w)

def build_workbook(data):
    wb = Workbook()
    ws = wb.active
    ws.title = '总览'
    headers = ['创建时间', '人员', '部门', '模板', '内容摘要']
    ws.append(headers)
    style_header(ws, len(headers), f'钉钉日志总览（{len(data)}条）')
    data_sorted = sorted(data, key=lambda x: x.get('create_time', 0))
    for x in data_sorted:
        text = content_text(x).replace('\n', ' / ')
        ws.append([
            fmt_ts(x.get('create_time', 0)),
            clean(x.get('creator_name', '')),
            clean(x.get('dept_name', '')),
            clean(x.get('template_name', '')),
            text[:300] + ('...' if len(text) > 300 else ''),
        ])
    auto_width(ws)

    # 各模板 sheet
    by_tpl = defaultdict(list)
    for x in data:
        by_tpl[x.get('template_name') or '未知'].append(x)

    for tpl, items in sorted(by_tpl.items(), key=lambda kv: -len(kv[1])):
        # 收集该模板所有出现的字段 key（保序）
        keys, seen = [], set()
        for x in sorted(items, key=lambda v: v.get('create_time', 0)):
            for c in (x.get('contents') or []):
                k = c.get('key')
                if k and k not in seen:
                    seen.add(k)
                    keys.append(k)
        sheet_name = tpl[:28] if len(tpl) <= 28 else tpl[:28]  # Excel sheet名限制31字符
        ws2 = wb.create_sheet(sheet_name)
        hdr = ['创建时间', '人员', '部门'] + keys + ['报告ID']
        ws2.append(hdr)
        style_header(ws2, len(hdr), f'{tpl}（{len(items)}条）')
        for x in sorted(items, key=lambda v: v.get('create_time', 0)):
            cmap = {c.get('key'): clean(c.get('value') or '') for c in (x.get('contents') or [])}
            row = [fmt_ts(x.get('create_time', 0)), clean(x.get('creator_name', '')), clean(x.get('dept_name', ''))]
            row += [cmap.get(k, '') for k in keys]
            row.append(clean(x.get('report_id', '')))
            ws2.append(row)
        auto_width(ws2)

    return wb

if __name__ == '__main__':
    data = json.loads(RAW.read_text(encoding='utf-8'))
    wb = build_workbook(data)
    wb.save(OUT)
    print(f'✅ 已导出: {OUT}')
    print(f'   共 {len(data)} 条, {len(wb.sheetnames)} 个 sheet: {wb.sheetnames}')
